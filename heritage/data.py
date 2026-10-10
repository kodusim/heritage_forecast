from __future__ import annotations
from pathlib import Path
import csv
import codecs
import re
import shutil
import zipfile
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from .storage import uid, sha

NS={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}

def csv_encoding(path):
    with Path(path).open('rb') as f:sample=f.read(65536)
    if sample.startswith((b'\xff\xfe',b'\xfe\xff')):return 'utf-16'
    for enc in ('utf-8-sig','cp949'):
        try: codecs.getincrementaldecoder(enc)().decode(sample,final=False); return enc
        except UnicodeError: pass
    raise ValueError('파일 인코딩을 인식할 수 없습니다. UTF-8 CSV로 저장해 주세요.')

def dates(series):
    if pd.api.types.is_datetime64_any_dtype(series):
        d=pd.to_datetime(series,errors='coerce')
    else:
        num=pd.to_numeric(series,errors='coerce')
        s=series.astype(str).str.strip()
        korean=s.str.contains('오전|오후',na=False)
        d=pd.to_datetime(s.where(num.isna()),format='mixed',errors='coerce')
        if korean.any():
            d.loc[korean]=pd.to_datetime(s[korean].str.replace('오전','AM').str.replace('오후','PM'),format='%Y-%m-%d %p %I:%M:%S',errors='coerce')
        excel=num.between(20000,80000)
        d.loc[excel]=pd.to_datetime(num[excel],unit='D',origin='1899-12-30').dt.round('s')
        compact=num.between(19550101,21000101)&num.eq(num.round())
        d.loc[compact]=pd.to_datetime(num[compact].astype('int64').astype(str),format='%Y%m%d',errors='coerce')
    try:
        if d.dt.tz is not None: d=d.dt.tz_convert('Asia/Seoul').dt.tz_localize(None)
    except AttributeError:
        raise ValueError('시간대가 혼합된 날짜입니다. 한 시간대로 통일해 주세요.')
    return d.where(d.between('1955-01-01','2100-01-01'))

def inspect_file(path):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        with path.open(encoding=csv_encoding(path),newline='') as f:
            rows=[]
            for _,r in zip(range(16),csv.reader(f)): rows.append(r)
        return [{'name':'CSV','rows':rows}]
    if path.suffix.lower()!='.xlsx': raise ValueError('지원 형식은 .xlsx와 .csv입니다.')
    try:
        w=load_workbook(path,read_only=True,data_only=True)
        sheets=[{'name':s.title,'rows':[list(r) for r in s.iter_rows(min_row=1,max_row=14,values_only=True)]} for s in w]
        w.close()
        return sheets
    except zipfile.BadZipFile as e:
        raise ValueError('손상되었거나 완전히 내려받지 않은 엑셀입니다.') from e

def infer_mapping(rows):
    width=max(map(len,rows),default=0)
    if not width: raise ValueError('빈 시트입니다.')
    padded=pd.DataFrame(rows).reindex(columns=range(width))
    scores=[int(dates(padded[c]).notna().sum()) for c in padded]
    time_col=int(np.argmax(scores))
    if scores[time_col]==0: raise ValueError('날짜 열을 찾지 못했습니다. 날짜 열과 헤더 행을 지정해 주세요.')
    parsed=dates(padded[time_col]); first=int(np.flatnonzero(parsed.notna())[0])
    header=max(0,first-1)
    for i,row in enumerate(rows[:first]):
        if any(isinstance(v,str) and re.fullmatch(r'(5|6|M)-[DT]\d(-[xy])?',v.strip()) for v in row):
            header=i;break
        if any(str(v).upper() in ['TIMESTAMP','DATE','일시','날짜'] for v in row): header=i
    return header,time_col

def mapping_names(rows,header,time_col):
    row=rows[header]; result={};seen={}
    parent=''
    for i,v in enumerate(row):
        if i==time_col: continue
        name=str(v).strip() if v is not None else ''
        if not name or name.lower() in ['none','nan']: name=f'열_{i+1}'
        if name.lower() in ['x','y'] and header>0:
            for j in range(i,-1,-1):
                above=rows[header-1][j] if j<len(rows[header-1]) else None
                if above is not None and str(above).strip(): parent=str(above).strip();break
            name=parent+'_'+name
        seen[name]=seen.get(name,0)+1
        if seen[name]>1: name+=f'_{i+1}'
        result[i]=name
    return result

def read_sheet(path,sheet,header,time_col,names,progress=lambda _:None,cancel=lambda:None):
    path=Path(path)
    if path.suffix.lower()=='.csv':
        raw=pd.read_csv(path,encoding=csv_encoding(path),header=None,skiprows=header+1,low_memory=False)
    else:
        wb=load_workbook(path,read_only=True,data_only=True)
        spec=wb[sheet]; xmlpath=spec._worksheet_path.lstrip('/');epoch=wb.epoch;wb.close()
        rows=[];strings=[];maxcol=max([time_col]+list(map(int,names)))
        with zipfile.ZipFile(path) as z:
            if 'xl/sharedStrings.xml' in z.namelist():
                with z.open('xl/sharedStrings.xml') as f:
                    for _,e in ET.iterparse(f,events=['end']):
                        if e.tag.endswith('}si'):
                            strings.append(''.join(t.text or '' for t in e.iter() if t.tag.endswith('}t')));e.clear()
            with z.open(xmlpath) as f:
                for _,e in ET.iterparse(f,events=['end']):
                    if not e.tag.endswith('}row'): continue
                    number=int(e.attrib.get('r',0))
                    if number<=header+1: e.clear();continue
                    vals={}
                    for c in e:
                        ref=c.attrib.get('r','A1');col=0
                        for ch in ref:
                            if not ch.isalpha(): break
                            col=col*26+ord(ch.upper())-64
                        col-=1
                        if col>maxcol: continue
                        node=c.find('m:v',NS);v=node.text if node is not None else None
                        typ=c.attrib.get('t')
                        if typ=='s' and v is not None: v=strings[int(v)]
                        elif typ=='inlineStr': v=''.join(t.text or '' for t in c.iter() if t.tag.endswith('}t'))
                        elif v is not None and typ not in ['str','e']:
                            try:v=float(v)
                            except ValueError:pass
                        if epoch.year==1904 and col==time_col and isinstance(v,(int,float)):v+=1462
                        vals[col]=v
                    if vals.get(time_col) is not None: rows.append([vals.get(i) for i in range(maxcol+1)])
                    e.clear()
                    if number%10000==0: cancel();progress(f'{sheet}: {number:,}행 읽는 중')
        raw=pd.DataFrame(rows)
    if raw.empty or time_col not in raw: raise ValueError('선택한 시트와 날짜 열에 자료가 없습니다.')
    dt=dates(raw[time_col]); valid=dt.notna()
    out=pd.DataFrame(index=pd.DatetimeIndex(dt[valid]));out.index.name='time'
    invalid=int((~valid).sum())
    for col,name in names.items():
        col=int(col)
        if col not in raw or not name.strip(): continue
        s=pd.to_numeric(raw.loc[valid,col],errors='coerce').replace([np.inf,-np.inf],np.nan)
        out[name]=s.to_numpy()
    if not out.shape[1]: raise ValueError('학습할 수치 열을 하나 이상 선택하세요.')
    return out.sort_index(kind='stable'),invalid

def merge_frames(old,new,policy='reject'):
    if new.columns.duplicated().any(): raise ValueError('센서 이름이 중복됩니다.')
    both=pd.concat([old,new]) if old is not None else new.copy()
    g=both.groupby(level=0,sort=True);counts=g.nunique(); conflict=counts>1
    n=int(conflict.sum().sum())
    if n and policy=='reject':
        raise ValueError(f'같은 시간·센서에 서로 다른 값이 {n:,}개 있습니다. 기존값 유지 / 새 값 사용 / 충돌값 제외 중 처리 방식을 선택하세요.')
    result=g.last() if policy=='new' else g.first()
    if policy=='mask': result=result.mask(conflict)
    return result,{'conflicts':n,'duplicate_rows':len(both)-len(result)}

def quality(df):
    rows=[]
    for col in df:
        s=df[col];v=s.dropna()
        rows.append({'sensor':col,'count':int(s.count()),'missing_pct':round(100*float(s.isna().mean()),3),
                     'first':str(s.first_valid_index()) if len(v) else '', 'last':str(s.last_valid_index()) if len(v) else '',
                     'min':float(v.min()) if len(v) else None,'max':float(v.max()) if len(v) else None})
    return rows

def import_data(project,req,progress=lambda _:None,cancel=lambda:None):
    path=Path(req['path']); digest=sha(path)
    prior=project.latest_dataset()
    signature=sha_mapping(req)
    from .history import disabled
    excluded = disabled(project)
    for rec in project.list('dataset'):
        if rec['id'] in excluded: continue
        if rec.get('source_hash')==digest and rec.get('mapping_hash')==signature:
            raise ValueError('동일한 파일과 매핑을 이미 등록했습니다. 데이터가 중복 추가되지 않았습니다.')
    parts=[];invalid=0
    for m in req['mappings']:
        frame,n=read_sheet(path,m['sheet'],m['header'],m['time_col'],m['columns'],progress,cancel)
        parts.append(frame);invalid+=n
    new=pd.concat(parts);cancel()
    old=project.load_data()[0] if prior else None
    df,issues=merge_frames(old,new,req.get('conflicts','reject'))
    if df.empty: raise ValueError('유효한 시각이 없습니다.')
    if len(df)>5_000_000: raise ValueError('프로젝트당 500만 시각을 초과합니다. 대상을 분리하세요.')
    meta=dict(prior.get('columns',{}) if prior else {})
    for name,description in req['columns_meta'].items():
        if name in meta and dict(meta[name],aggregation=meta[name].get('aggregation','mean'))!=dict(description,aggregation=description.get('aggregation','mean')):
            raise ValueError(f'{name}: 기존 역할·단위와 다릅니다. 센서 교체나 단위 변경은 새 센서 이름으로 등록하세요.')
        meta[name]=description
    ident=uid();rel=f'datasets/{ident}.csv.gz';dest=project.file(rel);dest.parent.mkdir(parents=True,exist_ok=True)
    temp=dest.with_name(dest.name+'.tmp')
    progress('데이터 버전과 품질 요약 저장 중');df.to_csv(temp,compression='gzip',float_format='%.17g');cancel();temp.replace(dest)
    src=project.file(f'sources/{digest}{path.suffix.lower()}');src.parent.mkdir(parents=True,exist_ok=True)
    if sha(path)!=digest: raise ValueError('읽는 동안 원본 파일이 변경되었습니다. 파일을 닫고 다시 등록하세요.')
    if not src.exists(): shutil.copy2(path,src)
    if sha(src)!=digest: raise ValueError('보관한 원본의 검증에 실패했습니다. 다시 등록하세요.')
    payload={'import_order':1+max((r.get('import_order',0) for r in project.list('dataset')),default=0),'file':rel,'parent':prior['id'] if prior else None,'rows':len(df),'start':str(df.index.min()),'end':str(df.index.max()),
             'columns':meta,'quality':quality(df),'source_hash':digest,'mapping_hash':signature,'source_name':path.name,
             'source_file':src.relative_to(project.path).as_posix(),'mappings':req['mappings'],'conflict_policy':req.get('conflicts','reject'),
             'invalid_dates':invalid,**issues}
    project.add('dataset',payload,ident)
    project.set_setting('active_dataset',ident)
    return ident

def sha_mapping(req):
    import hashlib,json
    return hashlib.sha256(json.dumps({k:req[k] for k in ('mappings','columns_meta','conflicts')},sort_keys=True).encode()).hexdigest()

def hourly(df,rule='1h',minimum=1,metadata=None):
    if df.index.duplicated().any(): raise ValueError('중복시각을 해결한 후 학습하세요.')
    if df.empty: raise ValueError('자료가 없습니다.')
    step=pd.Timedelta(rule)
    if (df.index.max()-df.index.min())/step>2_000_000: raise ValueError('시간 범위가 너무 큽니다. 날짜 매핑을 확인하세요.')
    agg=df.resample(rule,label='right',closed='left')
    result=agg.mean()
    for col in df:
        method=(metadata or {}).get(col,{}).get('aggregation','mean')
        series=df[col].resample(rule,label='right',closed='left')
        if method=='last':result[col]=series.last()
        elif method=='sum':result[col]=series.sum(min_count=1)
        elif method=='circular':
            angles=np.deg2rad(df[col]);args={'label':'right','closed':'left'}
            sine=np.sin(angles).resample(rule,**args).mean();cosine=np.cos(angles).resample(rule,**args).mean()
            result[col]=(np.rad2deg(np.arctan2(sine,cosine))%360).where(np.hypot(sine,cosine)>1e-8)
        elif method!='mean':raise ValueError(f'{col}: 지원하지 않는 시간 집계 방식입니다.')
    return result.where(agg.count()>=minimum)

def window_indices(frame,target,inputs,lookback,horizons,train_end,val_end):
    if target not in frame or not inputs: raise ValueError('대상·입력 센서를 선택하세요.')
    values=frame[inputs].to_numpy(dtype=float)
    good=np.isfinite(values).all(axis=1)
    rolling=pd.Series(good).rolling(lookback,min_periods=lookback).sum().eq(lookback).to_numpy()
    y=frame[target].to_numpy(dtype=float)
    idx=np.arange(lookback-1,len(frame)-max(horizons))
    if not len(idx): raise ValueError('입력기간과 예측시간에 비해 데이터가 부족합니다.')
    valid=rolling[idx]&np.isfinite(y[idx])
    for h in horizons: valid &= np.isfinite(y[idx+h])
    idx=idx[valid];t=frame.index
    tr=idx[t[idx+max(horizons)]<=pd.Timestamp(train_end)]
    va=idx[(t[idx+min(horizons)]>pd.Timestamp(train_end))&(t[idx+max(horizons)]<=pd.Timestamp(val_end))]
    te=idx[t[idx+min(horizons)]>pd.Timestamp(val_end)]
    return {'train':tr,'validation':va,'test':te}
