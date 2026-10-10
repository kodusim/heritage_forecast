"""Guided workflow: inspect measurements, train, inspect results, forecast a full path."""
from pathlib import Path
import pandas as pd
import numpy as np
from PySide6.QtCore import QThread,Signal,Qt,QTimer
from PySide6.QtWidgets import (QWidget,QVBoxLayout,QHBoxLayout,QFormLayout,QGridLayout,QTabWidget,
    QScrollArea,QSpinBox,QDoubleSpinBox,QComboBox,QCheckBox,QTableView,QSplitter,QFileDialog,QMessageBox,QDialog)
from .studio import MainWindow as Studio,DateField,MODEL,MODE,fold
from .ui import label,button,combo,set_table,FrameModel
from .storage import read_json
from .workflow import forecast_view


class Reader(QThread):
    ready=Signal(object);failed=Signal(str)
    def __init__(self,fn):super().__init__();self.fn=fn
    def run(self):
        try:self.ready.emit(self.fn())
        except Exception as e:self.failed.emit(str(e))


class MeasurementModel(FrameModel):
    def data(self,index,role=Qt.ItemDataRole.DisplayRole):
        if index.isValid() and role==Qt.ItemDataRole.DisplayRole:
            value=self.df.iloc[index.row(),index.column()]
            return '' if pd.isna(value) else str(value)
        return super().data(index,role)


class MainWindow(Studio):
    def __init__(self):
        self.readers=[];self.measurements=None;self.browse_key=None;self.checked_config=None;self.pending_config=None
        super().__init__()
        self.setWindowTitle('Heritage Studio · 데이터에서 미래 예측까지')
        self.job_label.parentWidget().layout().setContentsMargins(10,5,10,5)
        for index,text in [(0,'사용 순서'),(1,'1  데이터 넣고 확인'),(2,'2  모델 학습하기'),(5,'3  학습 결과 보기'),(3,'4  미래 움직임 예측'),(4,'5  논문 그림 저장')]:self.nav[index].setText(text)
        self.navigate(1)

    def page(self,title,description):
        if title=='새 예측':description='학습한 모델과 기간을 선택하면, 해당 센서의 마지막 관측 뒤에 이어지는 미래 곡선을 만듭니다.'
        return super().page(title,description)

    def background(self,fn,done):
        task=Reader(fn);self.readers.append(task);task.ready.connect(done);task.failed.connect(self.read_error)
        task.finished.connect(lambda:self.readers.remove(task) if task in self.readers else None);task.start()

    def read_error(self,error):
        self.job_label.setText('자료를 불러오지 못했습니다: '+error)

    def build_data(self):
        super().build_data();page=self.stack.widget(1);layout=page.layout();tabs=page.findChild(QTabWidget)
        browser=QWidget();v=QVBoxLayout(browser);row=QHBoxLayout();self.browse_sensor=QComboBox();self.browse_sensor.currentIndexChanged.connect(self.render_measurements)
        row.addWidget(label('센서'));row.addWidget(self.browse_sensor,1);row.addWidget(button('전체 기간',self.browse_all));v.addLayout(row)
        row=QHBoxLayout();self.browse_start=DateField();self.browse_end=DateField();row.addWidget(label('시작'));row.addWidget(self.browse_start);row.addWidget(label('종료'));row.addWidget(self.browse_end);row.addWidget(button('기간 적용',self.render_measurements));v.addLayout(row)
        self.browse_hint=label('파일을 추가하면 계측 표와 그래프가 여기에 표시됩니다.','muted');v.addWidget(self.browse_hint)
        views=QTabWidget();self.raw_table=QTableView();views.addTab(self.raw_table,'계측 표')
        self.raw_plot=QWidget();self.raw_plot_layout=QVBoxLayout(self.raw_plot);views.addTab(self.raw_plot,'센서 그래프');v.addWidget(views,1)
        row=QHBoxLayout();row.addWidget(button('현재 표 CSV 저장',self.export_measurements));row.addWidget(button('다음 → 모델 학습하기',lambda:self.navigate(2),True));v.addLayout(row)
        tabs.insertTab(0,browser,'계측 표 · 그래프');tabs.setCurrentIndex(0);self.data_tabs=tabs
        layout.itemAt(layout.count()-1).widget().hide()
        history_tab=tabs.widget(2);history_tab.layout().addWidget(button('선택 파일의 원본 표 열기',self.open_original))

    def refresh(self):
        old=getattr(self,'browse_key',None);super().refresh()
        if not hasattr(self,'browse_sensor'):return
        if hasattr(self,'future_run'):
            previous=self.future_run.currentData();self.future_run.blockSignals(True);self.future_run.clear()
            for run in self.project.list('run') if self.project else []:
                if run['status']=='complete':self.future_run.addItem(run['created'][:16]+' · '+', '.join(MODEL.get(k,k) for k in run['config']['models']),run['id'])
            self.future_run.setCurrentIndex(max(0,self.future_run.findData(previous)));self.future_run.blockSignals(False);self.filter_models()
        d=self.project.latest_dataset() if self.project else None;key=(str(self.project.path),d['id']) if d else None
        if old==key:return
        self.browse_key=key;self.measurements=None;self.checked_config=None
        if not d:
            self.browse_sensor.clear();set_table(self.raw_table,pd.DataFrame());self.clear_plot(self.raw_plot_layout)
            self.browse_hint.setText('현재 등록된 데이터가 없습니다. 엑셀 / CSV 추가부터 시작하세요.');return
        project=self.project;self.browse_hint.setText('계측 데이터를 불러오는 중입니다…')
        def loaded(result):
            if self.browse_key!=key:return
            frame,meta=result;self.measurements=frame;self.browse_sensor.blockSignals(True);self.browse_sensor.clear();self.browse_sensor.addItem('모든 센서 (표)','all')
            for name in frame:self.browse_sensor.addItem(name+' ['+meta['columns'][name]['unit']+']',name)
            self.browse_sensor.setCurrentIndex(1 if len(frame.columns) else 0);self.browse_sensor.blockSignals(False)
            self.browse_start.setText(str(max(frame.index.min(),frame.index.max()-pd.Timedelta(days=30))));self.browse_end.setText(str(frame.index.max()));self.render_measurements()
        self.background(lambda:project.load_data(d['id']),loaded)

    def browse_all(self):
        if self.measurements is None:return
        self.browse_start.setText(str(self.measurements.index.min()));self.browse_end.setText(str(self.measurements.index.max()));self.render_measurements()

    @staticmethod
    def clear_plot(layout):
        while layout.count():
            item=layout.takeAt(0)
            if item.widget():item.widget().deleteLater()

    def render_measurements(self,*args):
        if self.measurements is None:return
        start=pd.Timestamp(self.browse_start.text());end=pd.Timestamp(self.browse_end.text())
        if start>end:return self.browse_hint.setText('시작 날짜가 종료 날짜보다 늦습니다.')
        data=self.measurements.loc[start:end];sensor=self.browse_sensor.currentData()
        view=data if sensor=='all' else data[[sensor]] if sensor in data else data.iloc[:0]
        table=view.reset_index().rename(columns={view.index.name or 'index':'시각'})
        set_table(self.raw_table,table);self.raw_table.setModel(MeasurementModel(table,self.raw_table));self.clear_plot(self.raw_plot_layout)
        self.browse_hint.setText(f'선택 기간 {len(view):,}행 / 전체 {len(self.measurements):,}행. 등록 시 정렬·중복 처리한 계측값입니다. 원본 파일은 등록 내역에서 별도로 열 수 있습니다.')
        if sensor=='all':self.raw_plot_layout.addWidget(label('위에서 센서를 선택하면 그 센서의 계측 그래프가 표시됩니다.'));return
        if view.empty:self.raw_plot_layout.addWidget(label('선택 기간의 데이터가 없습니다.'));return
        from matplotlib.figure import Figure
        fig=Figure(layout='constrained');ax=fig.add_subplot();series=view[sensor]
        # Keep actual points, never interpolate; large ranges remain available in the table.
        if len(series)>100000:
            self.raw_plot_layout.addWidget(label('그래프는 100,000행 이하의 기간을 선택해 주세요. 표에는 선택한 전체 데이터를 표시했습니다.'));return
        ax.plot(series.index,series.values,lw=.8);ax.set_title(sensor);ax.set_ylabel(self.project.latest_dataset()['columns'][sensor]['unit']);ax.grid(alpha=.2)
        self.add_canvas(fig,self.raw_plot_layout)

    def export_measurements(self):
        if self.raw_table.model():self.save_csv(self.raw_table.model().df,'measurements.csv')

    def open_original(self):
        try:
            rec=self.selected(self.import_table,self.import_rows);path=self.project.file(rec['source_file'])
            dialog=QDialog(self);dialog.setWindowTitle('원본 표 · '+rec['source_name']);dialog.resize(1100,760);v=QVBoxLayout(dialog)
            sheet=QComboBox();sheets=list(dict.fromkeys(m['sheet'] for m in rec.get('mappings',[])))
            for name in sheets:sheet.addItem(name)
            page=QSpinBox();page.setRange(1,100000);row=QHBoxLayout();row.addWidget(label('등록한 시트'));row.addWidget(sheet,1);row.addWidget(label('페이지 (500행)'));row.addWidget(page);v.addLayout(row)
            info=label('');v.addWidget(info);table=QTableView();v.addWidget(table,1)
            from .data import csv_encoding
            def load():
                start=(page.value()-1)*500;name=sheet.currentText();info.setText('원본을 읽는 중…');page.setEnabled(False);sheet.setEnabled(False)
                def read():
                    if path.suffix.lower()=='.csv':return pd.read_csv(path,header=None,skiprows=start,nrows=500,encoding=csv_encoding(path),dtype=str,keep_default_na=False)
                    from openpyxl import load_workbook
                    book=load_workbook(path,read_only=True,data_only=True)
                    try:return pd.DataFrame(list(book[name].iter_rows(min_row=start+1,max_row=start+500,values_only=True)))
                    finally:book.close()
                def done(df):
                    set_table(table,df.astype(str).replace({'None':'','nan':''}));info.setText(f'원본 {start+1:,}행부터 {len(df):,}행 · 수식은 엑셀에 저장된 계산값을 표시합니다.');page.setEnabled(True);sheet.setEnabled(True)
                self.background(read,done)
            row.addWidget(button('페이지 보기',load));load();dialog.exec()
        except Exception as e:self.error(e)

    def build_training(self):
        super().build_training();old=self.stack.widget(2);self.stack.removeWidget(old)
        host=QWidget();layout=QVBoxLayout(host);layout.addWidget(label('모델 학습하기','pageTitle'));layout.addWidget(label('한 단계씩 선택하세요. 학습이 끝나면 오차·R²와 실제값 비교 그래프로 이어집니다.','muted'))
        tabs=QTabWidget();self.training_steps=tabs;layout.addWidget(tabs,1)
        def step(title):
            p=QWidget();v=QVBoxLayout(p);s=QScrollArea();s.setWidgetResizable(True);s.setWidget(p);tabs.addTab(s,title);return v
        v=step('1 센서');v.addWidget(label('어떤 센서의 움직임을 예측할까요?'));v.addWidget(self.targets);row=QHBoxLayout();row.addWidget(button('모두 선택',lambda:self.select_sensors(True)));row.addWidget(button('선택 해제',lambda:self.select_sensors(False)));v.addLayout(row);v.addStretch()
        v=step('2 예측 기간');v.addWidget(label('얼마 동안의 미래 곡선을 학습할까요?'))
        self.forecast_duration=QSpinBox();self.forecast_duration.setRange(1,365);self.forecast_duration.setValue(30)
        self.duration_unit=combo([('일',24),('시간',1)]);self.output_step=combo([('하루마다 예측값',24),('매시간 예측값',1)])
        form=QFormLayout();line=QWidget();row=QHBoxLayout(line);row.addWidget(self.forecast_duration);row.addWidget(self.duration_unit);form.addRow('예측할 기간',line);form.addRow('그래프의 예측 간격',self.output_step);v.addLayout(form)
        v.addWidget(label('예: 30일 · 하루마다 → 1일 뒤부터 30일 뒤까지 30개 값을 직접 학습합니다. 미래 환경값이나 가상의 관측을 만들어 넣지 않습니다.','muted'));v.addStretch()
        v=step('3 모델');v.addWidget(label('하나만 선택해도 되고, 여러 모델을 같은 조건으로 비교해도 됩니다.'))
        for q in self.model_checks.values():v.addWidget(q)
        self.model_checks['transformer'].setChecked(True);self.model_checks['lstm'].setChecked(False)
        v.addWidget(self.env_mode);v.addWidget(self.envs);v.addWidget(self.self_mode);v.addStretch()
        v=step('4 학습 설정');form=QFormLayout();self.learning_rate=QDoubleSpinBox();self.learning_rate.setDecimals(6);self.learning_rate.setRange(.000001,.1);self.learning_rate.setValue(.001);self.learning_rate.setSingleStep(.0001)
        self.model_width=combo([('작게 · 16',16),('기본 · 32',32),('크게 · 64',64),('매우 크게 · 128',128)]);self.model_width.setCurrentIndex(1)
        self.lookback.setValue(168);self.minimum.setValue(1);self.layer_count=QSpinBox();self.layer_count.setRange(1,4);self.dropout=QDoubleSpinBox();self.dropout.setRange(0,.8);self.dropout.setSingleStep(.05);self.dropout.setValue(.1)
        self.batch=combo([('32',32),('64',64),('128',128),('256',256)]);self.batch.setCurrentIndex(2)
        for title,w in [('과거 참고 기간 (시간)',self.lookback),('최대 학습 반복 횟수',self.epochs),('학습률 · 한 번에 조정하는 크기',self.learning_rate),('모델 크기',self.model_width)]:form.addRow(title,w)
        v.addLayout(form);advanced,av=fold('추가 설정');form=QFormLayout()
        for title,w in [('모델 층 수',self.layer_count),('드롭아웃',self.dropout),('배치 크기',self.batch),('시간당 최소 관측 수',self.minimum),('난수 값',self.seed),('계산 장치',self.device)]:form.addRow(title,w)
        av.addLayout(form);v.addWidget(advanced);v.addStretch()
        v=step('5 확인 · 학습');form=QFormLayout();form.addRow('학습 자료의 마지막 날짜',self.train_end);form.addRow('검증 자료의 마지막 날짜',self.val_end);v.addLayout(form)
        v.addWidget(label('날짜는 자동 제안합니다. 앞부분으로 학습하고 다음 구간으로 모델을 선택하며, 남은 마지막 구간의 실제값으로 최종 성능을 확인합니다.','muted'))
        self.check_table=QTableView();v.addWidget(self.check_table,1)
        self.training_summary=label('');layout.addWidget(self.training_summary)
        row=QHBoxLayout();back=button('← 이전',lambda:tabs.setCurrentIndex(max(0,tabs.currentIndex()-1)));next_button=button('다음 →',lambda:tabs.setCurrentIndex(min(4,tabs.currentIndex()+1)),True);self.train_button=button('자료 확인 후 학습 준비',self.start_training,True)
        row.addWidget(back);row.addStretch();row.addWidget(next_button);row.addWidget(self.train_button);layout.addLayout(row)
        def select(i):back.setEnabled(i>0);next_button.setVisible(i<4);self.train_button.setVisible(i==4);self.update_training_summary()
        tabs.currentChanged.connect(select);select(0);self.stack.insertWidget(2,host)
        # Legacy controls remain hidden for backwards-compatible result/figure methods.
        old.hide();old.setParent(host)
        for w in (self.forecast_duration,self.duration_unit,self.output_step):
            (w.valueChanged if hasattr(w,'valueChanged') else w.currentIndexChanged).connect(self.update_training_summary)

    def training_config(self):
        if not self.project or not self.project.latest_dataset():raise ValueError('먼저 데이터를 추가하세요.')
        total=self.forecast_duration.value()*self.duration_unit.currentData();step=self.output_step.currentData()
        if total<step or total%step:raise ValueError('하루보다 짧은 예측은 매시간 예측값을 선택하세요.')
        selected=[k for k,q in self.model_checks.items() if q.isChecked()]
        if not selected:raise ValueError('학습할 모델을 하나 이상 선택하세요.')
        return dict(dataset_id=self.project.latest_dataset()['id'],targets=self.checked(self.targets),environment=self.checked(self.envs),
            modes=(['self'] if self.self_mode.isChecked() else [])+(['environment'] if self.env_mode.isChecked() else []),models=selected,
            horizons=list(range(step,total+1,step)),forecast_step_hours=step,lookback=self.lookback.value(),epochs=self.epochs.value(),
            minimum=self.minimum.value(),seed=self.seed.value(),device=self.device.currentData(),train_end=self.train_end.text(),val_end=self.val_end.text(),
            learning_rate=self.learning_rate.value(),width=self.model_width.currentData(),layers=self.layer_count.value(),dropout=self.dropout.value(),batch_size=self.batch.currentData(),evaluation_stride=max(1,total//16))

    def update_training_summary(self,*args):
        if not hasattr(self,'forecast_duration'):return super().update_training_summary(*args)
        if not hasattr(self,'training_summary'):return
        total=self.forecast_duration.value()*self.duration_unit.currentData();step=self.output_step.currentData()
        self.training_summary.setText(f'센서 {len(self.checked(self.targets))}개 · {total/24:g}일 ({total}시간) · {step}시간 간격의 미래 곡선. 기존 학습 기록은 보존됩니다.')

    def start_training(self):
        if not self.idle():return
        try:
            cfg=self.training_config()
            if self.checked_config==cfg:self.start_job('train',cfg)
            else:self.pending_config=cfg;self.start_job('preflight',cfg)
        except Exception as e:self.error(e)

    def build_results(self):
        super().build_results()
        layout=self.stack.widget(5).layout();plot=self.result_tabs.widget(1)
        for i in (0,1):plot.layout().itemAt(i).widget().hide()
        actions=layout.itemAt(layout.count()-1).layout()
        actions.insertWidget(2,button('학습 설정 기록',self.show_experiments));actions.addWidget(button('다음 → 미래 예측',self.results_to_future,True))
        self.result_plot_timer=QTimer(self);self.result_plot_timer.setSingleShot(True);self.result_plot_timer.setInterval(300)
        self.result_plot_timer.timeout.connect(self.auto_comparison)

    def render_comparison(self,*args):
        super().render_comparison(*args)
        if hasattr(self,'result_summary'):
            df=self.comparison_data();rows=[]
            if not df.empty:
                r=self.project.get(self.run_combo.currentData());meta=self.project.get(r['config']['dataset_id'])['columns']
                for _,item in df.iterrows():
                    if item['model']=='persistence' and not self.show_baselines.isChecked():continue
                    rows.append({'센서':item.sensor,'모델':MODEL.get(item['model'],item['model']),'단위':meta[item.sensor]['unit'],
                        'MAE ↓':item.MAE,'RMSE ↓':item.RMSE,'R² ↑':item.R2,'평가 수':int(item.n)})
                from .history import affected
                warning='취소된 자료를 포함한 과거 결과입니다. ' if affected(self.project,r['config']['dataset_id']) else ''
                self.result_summary.setText(warning+f"{self.result_split.currentText()} · {self.result_horizon.currentText()}의 실제값과 비교한 성능입니다. 미래 예측의 성능을 보장하는 값은 아닙니다.")
            set_table(self.metrics_table,pd.DataFrame(rows))
        if hasattr(self,'result_plot_timer') and self.stack.currentIndex()==5 and not getattr(self,'drawing_comparison',False):self.result_plot_timer.start()

    def auto_comparison(self):
        if self.comparison_data().empty:return
        self.drawing_comparison=True
        try:self.comparison_plot()
        finally:self.drawing_comparison=False

    def comparison_plot(self):
        tab=self.result_tabs.currentIndex();super().comparison_plot();self.result_tabs.setCurrentIndex(tab)
        for i in range(self.compare_layout.count()):
            widget=self.compare_layout.itemAt(i).widget()
            if widget and hasattr(widget,'figure'):widget.setMinimumHeight(220)

    def show_experiments(self):
        if not self.project:return
        rows=[]
        for run in self.project.list('run'):
            cfg=run['config'];row={'학습일':run['created'][:19],'상태':run['status'],'모델':', '.join(MODEL.get(k,k) for k in cfg['models']),
                '예측 범위 (시간)':max(cfg['horizons']),'참고 기간 (시간)':cfg['lookback'],'학습률':cfg.get('learning_rate',.001),
                '모델 크기':cfg.get('width',32),'층 수':cfg.get('layers',1),'최대 반복':cfg['epochs'],'데이터 버전':cfg['dataset_id'][:8]}
            for mid in run['models']:
                m=self.project.get(mid)
                score=next((x for x in m['metrics'] if x['split']=='test' and x['horizon']==self.result_horizon.currentData()),{})
                rows.append(dict(row,모델=MODEL.get(m['model'],m['model']),센서=m['sensor'],입력=MODE[m['mode']],MAE=score.get('MAE'),RMSE=score.get('RMSE'),R2=score.get('R2'),평가간격=self.result_horizon.currentData(),학습종료=cfg['train_end'],검증종료=cfg['val_end']))
            if not run['models']:rows.append(row)
        dialog=QDialog(self);dialog.setWindowTitle('학습 기록과 설정');dialog.resize(1150,650);v=QVBoxLayout(dialog)
        v.addWidget(label('결과 화면의 학습 기록을 바꾸면 해당 설정의 MAE·RMSE·R²와 그래프를 확인할 수 있습니다. 데이터·평가 기간이 다른 기록은 직접적인 성능 비교 대상이 아닙니다.','muted'))
        table=QTableView();set_table(table,pd.DataFrame(rows));v.addWidget(table,1);dialog.exec()

    def results_to_future(self):
        rid=self.run_combo.currentData();self.navigate(3)
        if rid:
            self.future_run.setCurrentIndex(self.future_run.findData(rid))
            for i,m in enumerate(self.model_rows):
                if m['run_id']==rid and m['model']!='persistence':self.models_table.selectRow(i);break

    def build_models(self):
        super().build_models();fresh=self.predict_tabs.widget(0);layout=fresh.layout()
        # Historical replay remains available in a separate tab. The primary action
        # needs only a trained model and a duration after its latest observation.
        detail=QWidget();dv=QVBoxLayout(detail)
        while layout.count()>3:
            item=layout.takeAt(3)
            if item.widget():dv.addWidget(item.widget())
            elif item.layout():dv.addLayout(item.layout())
        self.predict_tabs.addTab(detail,'과거 시점 재현 · 상세')
        self.future_run=QComboBox();self.future_run.currentIndexChanged.connect(self.filter_models);layout.insertWidget(0,label('사용할 학습 기록'));layout.insertWidget(1,self.future_run)
        form=QFormLayout();self.future_duration=QSpinBox();self.future_duration.setRange(1,8760);self.future_duration.setValue(7)
        self.future_unit=combo([('일',24),('시간',1)]);line=QWidget();row=QHBoxLayout(line);row.addWidget(self.future_duration);row.addWidget(self.future_unit);form.addRow('마지막 관측 이후 예측할 기간',line);layout.addLayout(form)
        self.future_hint=label('기간 전체를 학습한 모델을 선택하세요. 예측은 해당 센서의 마지막 관측에서 시작합니다.','muted');layout.addWidget(self.future_hint)
        self.future_button=button('미래 움직임 예측하기',self.make_future,True);layout.addWidget(self.future_button)
        self.future_plot_tab=QWidget();v=QVBoxLayout(self.future_plot_tab);self.future_plot_hint=label('');v.addWidget(self.future_plot_hint);self.future_plot_host=QWidget();self.future_plot_layout=QVBoxLayout(self.future_plot_host);v.addWidget(self.future_plot_host,1)
        row=QHBoxLayout();row.addWidget(button('예측 표 저장',self.export_prediction));row.addWidget(button('그래프를 논문 그림으로 저장',self.save_future_plot));v.addLayout(row);self.predict_tabs.insertTab(1,self.future_plot_tab,'미래 예측 그래프')

    def filter_models(self,*args):
        models=getattr(self,'all_models',[])
        if hasattr(self,'future_run') and self.future_run.currentData():self.all_models=[m for m in models if m['run_id']==self.future_run.currentData()]
        try:super().filter_models(*args)
        finally:self.all_models=models
        if self.models_table.model() and self.model_rows:
            table=self.models_table.model().df.copy()
            for i,m in enumerate(self.model_rows):
                if m.get('forecast_step_hours'):table.loc[i,'예측 (시간 뒤)']=f"{max(m['horizons'])/24:g}일 / {m['forecast_step_hours']}시간 간격"
            set_table(self.models_table,table);self.models_table.selectionModel().selectionChanged.connect(self.show_model_detail);self.models_table.selectRow(0)

    def show_model_detail(self,*args):
        super().show_model_detail(*args)
        if not hasattr(self,'future_hint'):return
        try:
            m=self.selected(self.models_table,self.model_rows);step=m.get('forecast_step_hours')
            self.model_hint.setText(f"{m['sensor']} · {MODEL[m['model']]} · {MODE[m['mode']]} · 과거 {m['lookback']}시간 참고 · 단위 {m['columns_meta'][m['sensor']]['unit']}")
            self.future_hint.setText(f"{m['sensor']} · {MODEL[m['model']]} · 최대 {max(m['horizons'])/24:g}일, {step}시간 간격으로 학습한 모델입니다." if step else '이전 모델은 일부 시점만 예측합니다. 학습하기에서 기간 전체를 예측하는 모델을 만드세요.')
            from .history import affected
            valid=bool(step) and m['status']=='complete' and self.project.get(m['run_id'])['status']=='complete' and not affected(self.project,m['dataset_id'])
            self.future_button.setEnabled(valid)
            if affected(self.project,m['dataset_id']):self.future_hint.setText('취소된 자료로 학습한 모델입니다. 현재 데이터로 다시 학습하세요.')
            if step:
                unit=24 if step%24==0 else 1;self.future_unit.setCurrentIndex(0 if unit==24 else 1);self.future_duration.setValue(max(m['horizons'])//unit)
        except ValueError:self.future_button.setEnabled(False)

    def make_future(self):
        if not self.idle():return
        try:
            m=self.selected(self.models_table,self.model_rows);hours=self.future_duration.value()*self.future_unit.currentData()
            if hours>max(m['horizons']):raise ValueError('이 모델이 학습한 기간보다 깁니다. 학습하기의 예측 기간을 늘려 새 모델을 만들어 주세요.')
            self.start_job('future_curve',{'model_id':m['id'],'hours':hours})
        except Exception as e:self.error(e)

    def add_canvas(self,fig,layout):
        from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg,NavigationToolbar2QT
        from matplotlib.dates import AutoDateLocator,ConciseDateFormatter
        for ax in fig.axes:
            locator=AutoDateLocator(minticks=3,maxticks=6);ax.xaxis.set_major_locator(locator);ax.xaxis.set_major_formatter(ConciseDateFormatter(locator))
        canvas=FigureCanvasQTAgg(fig);canvas.setMinimumHeight(200);layout.addWidget(NavigationToolbar2QT(canvas,self));layout.addWidget(canvas,1);canvas.draw_idle();return canvas

    def show_future(self,ident):
        self.current_future=ident;rec,observed,pred=forecast_view(self.project,ident)
        from matplotlib.figure import Figure
        fig=Figure(layout='constrained');ax=fig.add_subplot();ax.plot(observed.index,observed.values,color='#334155',label='Observed',lw=1)
        future=pd.Series([observed.iloc[-1],*pred.predicted],index=[pd.Timestamp(rec['origin']),*pd.to_datetime(pred.time)])
        ax.plot(future.index,future.values,color='#087f70',label='Forecast',lw=1.5);ax.axvline(pd.Timestamp(rec['origin']),color='#999999',ls='--',lw=.8)
        ax.set_title(rec['sensor']);ax.set_ylabel(rec['unit']);ax.grid(alpha=.2);ax.legend();self.clear_plot(self.future_plot_layout);self.add_canvas(fig,self.future_plot_layout);self.future_figure=fig
        self.future_plot_hint.setText(f"최근 관측(시간 집계) → 미래 예측 · {rec['origin']} 이후 {rec.get('requested_hours',max(pred.horizon))}시간 · 실제 미래값이 없어 MAE·R²는 아직 계산할 수 없습니다.")
        from .history import affected
        if rec.get('historical_origin') and not rec.get('curve'):self.future_plot_hint.setText('과거 시점 재현입니다. '+self.future_plot_hint.text())
        if affected(self.project,rec['dataset_id']) or affected(self.project,self.project.get(rec['model_id'])['dataset_id']):self.future_plot_hint.setText('취소된 자료를 포함한 과거 기록입니다. '+self.future_plot_hint.text())
        self.predict_tabs.setCurrentWidget(self.future_plot_tab)
        index=next((i for i,p in enumerate(self.pred_rows) if p['id']==ident),None)
        if index is not None:self.predictions_table.selectRow(index)

    def view_prediction(self):
        try:self.show_future(self.selected(self.predictions_table,self.pred_rows)['id'])
        except Exception as e:self.error(e)

    def save_future_plot(self):
        if not hasattr(self,'future_figure'):return
        path,_=QFileDialog.getSaveFileName(self,'미래 예측 그림 저장','forecast.pdf','PDF (*.pdf);;SVG (*.svg);;PNG (*.png);;TIFF (*.tiff)')
        if path:
            original_size=self.future_figure.get_size_inches().copy()
            try:
                self.future_figure.set_size_inches(180/25.4,100/25.4)
                self.future_figure.savefig(path,dpi=300,facecolor='white')
            finally:self.future_figure.set_size_inches(original_size)
            rec,observed,pred=forecast_view(self.project,self.current_future);pred.to_csv(Path(path).with_suffix('.csv'),index=False,encoding='utf-8-sig')
            observed.rename('observed').to_csv(Path(path).with_name(Path(path).stem+'_observed.csv'),encoding='utf-8-sig')
            from .storage import write_json
            write_json(Path(path).with_suffix('.json'),dict(prediction_id=rec['id'],model_id=rec['model_id'],dataset_id=rec['dataset_id'],dpi=300,width_mm=180,height_mm=100))
            self.job_label.setText('그림과 예측 표, 모델·데이터 기록을 저장했습니다.')

    def job_finished(self,*args):
        kind=getattr(self,'active_kind','');p=self.job_path.with_suffix('.result.json');result=read_json(p) if p.exists() else {}
        super().job_finished(*args)
        if not result.get('ok'):return
        if kind=='preflight':
            rows=result['id'];set_table(self.check_table,pd.DataFrame(rows).rename(columns={'sensor':'센서','ready':'학습 가능','train':'학습 표본','validation':'검증 표본','test':'평가 표본','reason':'확인할 내용'}))
            if rows and all(r['ready'] for r in rows):self.checked_config=self.pending_config;self.train_button.setText('확인한 설정으로 학습 시작')
            else:self.checked_config=None;self.train_button.setText('설정을 바꾼 뒤 다시 확인');self.job_label.setText('선택한 센서 중 자료가 부족한 센서가 있습니다. 표의 안내를 확인하세요.')
        elif kind=='train':
            self.run_combo.setCurrentIndex(self.run_combo.findData(result['id']));self.show_metrics();self.result_horizon.setCurrentIndex(self.result_horizon.count()-1);self.auto_comparison();self.navigate(5)
        elif kind in ('future_curve','predict'):self.show_future(result['id']);self.navigate(3)
        elif kind=='import':self.data_tabs.setCurrentIndex(0);self.navigate(1)

    def navigate(self,n):
        super().navigate(n)
        if n==5 and hasattr(self,'result_plot_timer'):self.result_plot_timer.start()

    def closeEvent(self,event):
        if any(r.isRunning() for r in self.readers):
            self.job_label.setText('자료를 읽는 중입니다. 읽기가 끝나면 창을 닫을 수 있습니다.');event.ignore();return
        super().closeEvent(event)
